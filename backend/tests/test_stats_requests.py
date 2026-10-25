"""PR #10: "Yangi" belgisi (har xodim uchun), arizalar filtrlari va kartasi, statistika detali, Excel/PDF eksport."""
import io
from datetime import timedelta

from openpyxl import load_workbook
from sqlalchemy import update

from app.db.base import utcnow
from app.models import Request
from app.models.enums import RequestStatus
from tests.test_permissions import _manager_a


async def test_new_badge_is_per_user(world, client_factory):
    gps, admin = await client_factory("gps"), await client_factory("admin")
    rows = (await gps.get("/api/v1/gps/requests")).json()
    assert [r["id"] for r in rows] == sorted(world["requests"], reverse=True)  # yangisi tepada
    assert all(r["seen"] is False for r in rows)
    assert (await gps.post("/api/v1/requests/seen", json={"request_ids": [rows[0]["id"]]})).status_code == 204
    rows = (await gps.get("/api/v1/gps/requests")).json()
    assert [r["seen"] for r in rows] == [True, False, False]
    # Boshqa xodim uchun hali "yangi"
    page = (await admin.get("/api/v1/requests")).json()
    assert all(i["seen"] is False for i in page["items"])
    # Kartani ochish ham "ko'rildi" qiladi
    d = (await admin.get(f"/api/v1/requests/{rows[1]['id']}")).json()
    assert d["seen"] is True
    assert next(i for i in (await admin.get("/api/v1/requests")).json()["items"] if i["id"] == rows[1]["id"])["seen"] is True


async def test_request_detail_and_filters(world, client_factory, sessionmaker):
    admin, gps, driver = await client_factory("admin"), await client_factory("gps"), await client_factory("driver")
    from tests.helpers import open_and_start_shift
    await open_and_start_shift(gps, driver, world["users"]["driver"], world["vehicles"]["v1"])
    rid = world["requests"][0]
    await gps.post(f"/api/v1/gps/requests/{rid}/assign", json={"driver_id": world["users"]["driver"]})
    d = (await admin.get(f"/api/v1/requests/{rid}")).json()
    assert d["driver_name"] == "Haydovchi A" and d["naryad_no"] and [e["status"] for e in d["events"]] == ["assigned"]
    assert d["events"][0]["actor_name"] == "GPS"
    async with sessionmaker() as s:
        await s.execute(update(Request).where(Request.id == world["requests"][1]).values(is_repeat=True, waste_type="septic"))
        await s.commit()
    q = lambda **p: admin.get("/api/v1/requests", params=p)  # noqa: E731
    assert (await q(repeat="true")).json()["total"] == 1
    assert (await q(waste_type="septic")).json()["total"] == 1
    assert (await q(status="assigned")).json()["total"] == 1
    today = (utcnow() + timedelta(hours=5)).date()
    assert (await q(date_from=str(today), date_to=str(today))).json()["total"] == 3
    assert (await q(date_from=str(today + timedelta(days=1)))).json()["total"] == 0
    # Manager — o'z tumani; boshqa tuman arizasi kartasi yopiq
    _, mgr = await _manager_a(world, client_factory, admin, [])
    assert (await mgr.get("/api/v1/requests")).json()["total"] == 3
    async with sessionmaker() as s:
        r = Request(customer_id=world["customer"].id, district_id=world["districts"]["b"], status=RequestStatus.new)
        s.add(r); await s.commit()
        other = r.id
    assert (await mgr.get(f"/api/v1/requests/{other}")).status_code == 403


async def test_stats_detail_per_metric(world, client_factory, sessionmaker):
    admin = await client_factory("admin")
    r1, r2, r3 = world["requests"]
    now = utcnow()
    async with sessionmaker() as s:
        await s.execute(update(Request).where(Request.id == r1).values(status=RequestStatus.done, finished_at=now + timedelta(hours=30)))
        await s.execute(update(Request).where(Request.id == r2).values(due_at=now - timedelta(hours=1)))
        await s.execute(update(Request).where(Request.id == r3).values(is_repeat=True))
        await s.commit()
    get = lambda m, p="last5": admin.get("/api/v1/stats/detail", params={"period": p, "metric": m})  # noqa: E731
    d = (await get("received")).json()
    assert d["value"] == 3 and d["granularity"] == "day" and sum(x["value"] for x in d["series"]) == 3
    assert {c["key"]: c["value"] for c in d["composition"]} == {"done": 1, "pending": 1, "overdue": 1}
    assert {x["name"]: x["value"] for x in d["districts"]} == {"Asaka tumani": 3, "Termiz shahri": 0}
    d = (await get("done")).json()
    assert d["value"] == 1 and {c["key"]: c["value"] for c in d["composition"]} == {"on_time": 1, "late": 0, "suspect": 0}
    assert {h["bucket"]: h["value"] for h in d["histogram"]}["24-48"] == 1
    d = (await get("repeat")).json()
    assert d["value"] == 1 and sum(x["value"] for x in d["series"]) == 1  # faqat qayta arizalar — jami emas
    d = (await get("not_done")).json()
    assert d["value"] == 2 and {c["key"]: c["value"] for c in d["composition"]} == {"unassigned": 2, "in_progress": 0, "failed": 0}
    d = (await get("overdue", "today")).json()
    assert d["value"] == 1 and d["granularity"] == "hour"
    assert (await get("bogus")).status_code == 422


async def test_stats_export_xlsx_and_pdf(world, client_factory):
    admin = await client_factory("admin")
    r = await admin.get("/api/v1/stats/export", params={"period": "last10", "metric": "done", "format": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert 'filename="statistika-last10-done-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tumanlar", "Dinamika", "Tarkib", "Vaqt taqsimoti"]
    ws = wb["Tumanlar"]
    rows = [[v for v in (c.value for c in row)] for row in ws.iter_rows()]
    rows = [r[:max((i + 1 for i, v in enumerate(r) if v is not None), default=0)] for r in rows]  # oxiridagi bo'sh kataklarsiz
    assert rows[1] == ["Davr", "Oxirgi 10 kun"]
    assert ["Asaka tumani", 3, 0, 3, 0, 0, 3, 0, 100] in rows and rows[-1][0] == "Jami"
    assert wb["Tarkib"]["A1"].value == "Bajarilgan — tarkibi"
    r = await admin.get("/api/v1/stats/export", params={"period": "today", "metric": "received", "format": "pdf"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content[:4] == b"%PDF"
    import pdfplumber
    with pdfplumber.open(io.BytesIO(r.content)) as pdf:
        text = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
    assert "Arizalar statistikasi" in text and "Asaka tumani" in text and "Tushgan arizalar" in text
    assert (await (await client_factory("operator")).get("/api/v1/stats/export")).status_code == 403
