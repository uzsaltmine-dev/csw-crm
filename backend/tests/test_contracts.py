"""PR #17: «Yangi shartnoma» jarayoni — forma, tekshirish, raqam, hujjat (PDF/DOCX), imzolash, buxgalter tasdig'i/qaytarish, hisob-kitob."""
import io
from datetime import date, timedelta

from docx import Document
from sqlalchemy import update

from app.db.base import utcnow
from app.models import OrderForm, Request
from app.models.enums import RequestStatus
from app.services.contracts import months_between, totals
from app.models.enums import ContractType
from tests.test_accounting import _setup

T = date.today()
REQ = {"name": "\"Yangi Test\" MChJ", "stir": "309111222", "director": "Aliyev Vali", "phone": "90 123 45 67",
       "legal_address": "Asaka sh., Bog'bon ko'chasi 5", "address": "Asaka sh., Bog'bon ko'chasi 5", "bank": "DEMO bank",
       "mfo": "00873", "account": "20208000900123456001", "oked": "56290", "cadastre": "17:15:41:01:02:0001"}


def body(cid, **kw):
    b = {"customer_id": cid, "requisites": dict(REQ), "type": "by_volume", "price": 84000, "annual_volume": 100,
         "valid_from": str(T - timedelta(days=20)), "valid_to": str(T.replace(month=12, day=31) if T.month < 12 else T + timedelta(days=200)),
         "place": "Asaka tumani", "sign_method": "electronic"}
    b.update(kw)
    return b


async def _new_customer(legal, world) -> int:
    r = await legal.post("/api/v1/customers", json={"name": "Yangi Test MChJ", "district_id": world["districts"]["a"]})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_totals_vat_inside():
    t = totals(ContractType.by_volume, 84000, 142.92, None, None, None)
    assert str(t["total"]) == "12005280.00" and str(t["vat"]) == "1286280.00" and str(t["unit_net"]) == "75000.00"
    assert months_between(date(2026, 3, 1), date(2026, 12, 31)) == 10


async def test_full_flow(world, client_factory, sessionmaker):
    await _setup(world, sessionmaker)  # buxgalter foydalanuvchisi
    legal, acc, gps, op = [await client_factory(u) for u in ("legal", "buxgalter", "gps", "operator")]
    cid = await _new_customer(legal, world)
    # Tekshirish: rekvizitlar yetishmasa — maydonlar bo'yicha xatolar
    bad = body(cid, requisites={"name": "X MChJ", "stir": "12"}, annual_volume=None)
    r = (await legal.post("/api/v1/contracts/check", json=bad)).json()
    fields = {e["field"] for e in r["errors"]}
    assert not r["ok"] and {"stir", "director", "account", "mfo", "phone", "annual_volume"} <= fields
    r = (await legal.post("/api/v1/contracts/check", json=body(cid))).json()
    assert r["ok"] and r["totals"]["total"] == 8_400_000 and r["totals"]["words"].startswith("Sakkiz million")
    assert (await op.post("/api/v1/contracts/check", json=body(cid))).status_code == 403
    # Saqlash: ichki raqam viloyat bo'yicha, loyiha — faol emas
    r = await legal.post("/api/v1/contracts", json=body(cid))
    assert r.status_code == 201, r.text
    ct = r.json()
    assert ct["number"] == f"BTE-{(T - timedelta(days=20)).year}/00001" and ct["status"] == "draft" and not ct["is_active"]
    assert ct["customer"]["account"] == REQ["account"] and ct["customer"]["phone"] == "+998 90 123-45-67"
    # Hujjat: PDF (loyiha belgisi bilan) va DOCX
    pdf = await legal.get(f"/api/v1/contracts/{ct['id']}/document", params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    dx = await legal.get(f"/api/v1/contracts/{ct['id']}/document", params={"format": "docx"})
    text = "\n".join(p.text for p in Document(io.BytesIO(dx.content)).paragraphs)
    assert "OMMAVIY SHARTNOMA" in text and "Sakkiz million" in text and REQ["cadastre"] in text
    # Ikkinchi loyiha ustma-ust tushadi — faqat imzolangandan keyin; loyiha ustma-ust hisoblanmaydi
    # Imzolash: elektron — platforma va ID majburiy
    sign = {"signed_on": str(T), "sign_method": "electronic", "platform": "didox"}
    r = await legal.post(f"/api/v1/contracts/{ct['id']}/sign", data=sign)
    assert r.json()["error"]["code"] == "platform_required"
    r = await legal.post(f"/api/v1/contracts/{ct['id']}/sign", data={**sign, "platform_id": "e0f94e102ced"},
                         files={"file": ("imzolangan.pdf", b"%PDF-1.4 test", "application/pdf")})
    assert r.status_code == 200, r.text
    ct = r.json()
    assert ct["status"] == "signed" and ct["is_active"] and ct["display_number"] == "e0f94e102ced" and ct["has_scan"]
    assert (await legal.get(f"/api/v1/contracts/{ct['id']}/scan")).content == b"%PDF-1.4 test"
    # Imzolangan shartnoma o'zgartirilmaydi; yangi ustma-ust shartnoma — xato
    assert (await legal.put(f"/api/v1/contracts/{ct['id']}", json=body(cid))).json()["error"]["code"] == "contract_locked"
    r = (await legal.post("/api/v1/contracts/check", json=body(cid))).json()
    assert {"field": "valid_from", "code": "overlap"} in r["errors"]
    # Korxona endi faol: GPS kartasida «Yangi shartnoma»
    c = (await legal.get(f"/api/v1/customers/{cid}")).json()
    assert c["missing_fields"] == [] and c["new_contract"] is True
    # Buxgalter navbati; legal tasdiqlay olmaydi
    q = (await acc.get("/api/v1/contracts", params={"status": "signed"})).json()
    assert any(x["id"] == ct["id"] for x in q["items"])
    assert (await legal.post(f"/api/v1/contracts/{ct['id']}/approve")).status_code == 403
    assert (await gps.get("/api/v1/contracts")).status_code == 403
    # Qaytarish (izoh bilan) → legal tuzatadi → qayta yuboradi → tasdiq
    r = (await acc.post(f"/api/v1/contracts/{ct['id']}/return", json={"comment": "Narx noto'g'ri"})).json()
    assert r["status"] == "returned" and r["return_comment"] == "Narx noto'g'ri" and r["is_active"]
    r = await legal.put(f"/api/v1/contracts/{ct['id']}", json=body(cid, price=90000))
    assert r.status_code == 200 and r.json()["price"] == 90000
    r = (await legal.post(f"/api/v1/contracts/{ct['id']}/sign", data={**sign, "platform_id": "e0f94e102ced"})).json()
    assert r["status"] == "signed" and r["return_comment"] is None
    r = (await acc.post(f"/api/v1/contracts/{ct['id']}/approve")).json()
    assert r["status"] == "approved"
    assert [e["action"] for e in r["events"]] == ["created", "signed", "returned", "edited", "resubmitted", "approved"]


async def test_pending_not_billed_until_approved(world, client_factory, sessionmaker):
    await _setup(world, sessionmaker)
    legal, acc = await client_factory("legal"), await client_factory("buxgalter")
    cid = await _new_customer(legal, world)
    ct = (await legal.post("/api/v1/contracts", json=body(cid, sign_method="paper"))).json()
    ct = (await legal.post(f"/api/v1/contracts/{ct['id']}/sign", data={"signed_on": str(T), "sign_method": "paper"})).json()
    assert ct["status"] == "signed" and ct["platform_id"] is None and not ct["has_scan"]
    # Shu korxonaga bajarilgan murojaat (5 m³)
    rid = world["requests"][2]
    async with sessionmaker() as s:
        await s.execute(update(Request).where(Request.id == rid).values(customer_id=cid, contract_id=ct["id"], status=RequestStatus.done,
                                                                        finished_at=utcnow() - timedelta(days=1)))
        s.add(OrderForm(request_id=rid, number="7777777", print_count=1, volume_m3=5, waste_kind="solid",
                        closed_at=utcnow() - timedelta(days=1), closed_via="form"))
        await s.commit()
    month = {"month": str((T - timedelta(days=1)).replace(day=1))}
    f = (await acc.get("/api/v1/accounting/order-forms", params=month)).json()
    item = next(x for x in f["items"] if x["customer_id"] == cid)
    assert item["pending"] and item["amount"] == 0 and f["pending"] >= 1
    g = (await acc.post("/api/v1/accounting/invoices/generate", params=month)).json()
    assert g["pending_forms"] >= 1
    invs = (await acc.get("/api/v1/accounting/invoices", params=month)).json()["items"]
    assert not any(i["customer_id"] == cid for i in invs)
    row = next((x for x in (await acc.get("/api/v1/accounting/companies")).json()["items"] if x["customer_id"] == cid), None)
    assert row is not None and row["debit"] == 0
    # Tasdiqdan keyin — hisobga kiradi (5 m³ × 84 000)
    await acc.post(f"/api/v1/contracts/{ct['id']}/approve")
    row = next(x for x in (await acc.get("/api/v1/accounting/companies")).json()["items"] if x["customer_id"] == cid)
    assert row["debit"] == 420_000
    await acc.post("/api/v1/accounting/invoices/generate", params=month)
    inv = next(i for i in (await acc.get("/api/v1/accounting/invoices", params=month)).json()["items"] if i["customer_id"] == cid)
    assert inv["total"] == 420_000
    lst = (await acc.get("/api/v1/contracts", params={"q": "Yangi"})).json()["items"]
    assert lst[0]["used_m3"] == 5 and lst[0]["usage_pct"] == 5.0


async def test_draft_delete_and_scope(world, client_factory):
    legal = await client_factory("legal")
    cid = await _new_customer(legal, world)
    ct = (await legal.post("/api/v1/contracts", json=body(cid))).json()
    assert (await legal.delete(f"/api/v1/contracts/{ct['id']}")).status_code == 204
    assert (await legal.get(f"/api/v1/contracts/{ct['id']}")).status_code == 404
    # Surxondaryo (BTK) — Didox shabloni
    r = await legal.post("/api/v1/customers", json={"name": "Termiz Test MChJ", "district_id": world["districts"]["b"]})
    ct = (await legal.post("/api/v1/contracts", json=body(r.json()["id"], requisites={**REQ, "stir": "309111333"}))).json()
    assert ct["number"].startswith("BTK-")
    dx = await legal.get(f"/api/v1/contracts/{ct['id']}/document", params={"format": "docx"})
    text = "\n".join(p.text for p in Document(io.BytesIO(dx.content)).paragraphs)
    assert "Шартнома №" in text and "Sakkiz million" in text
