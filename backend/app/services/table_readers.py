"""Fayldan jadval qatorlarini o'qish: .xlsx, .xls, .csv, .docx, .pdf (matnli).

Har format bir xil ko'rinishga keltiriladi — `Table(title, rows)`, qatorlar katakchalar kortejlari.
Keyingi qadamlar (sarlavhani topish, tumanlarni bog'lash, birlashtirish) formatdan qat'i nazar bir xil.
Skanerlangan PDF (matnsiz) — `pdf_scanned` xatosi: OCR keyingi bosqichda shu joyga ulanadi.
"""
import csv
import io
from collections.abc import Iterable, Iterator
from dataclasses import dataclass

from app.core.errors import invalid

SUPPORTED = (".xlsx", ".xls", ".csv", ".docx", ".pdf")


@dataclass
class Table:
    title: str
    rows: Iterable[tuple]


def file_kind(file_name: str) -> str:
    name = (file_name or "").lower()
    for ext in SUPPORTED:
        if name.endswith(ext):
            return ext[1:]
    raise invalid("unsupported_format", "Qo'llab-quvvatlanadigan formatlar: " + ", ".join(SUPPORTED))


def _bad(kind: str) -> Exception:
    return invalid("bad_file", f"Faylni o'qib bo'lmadi: fayl buzilgan yoki .{kind} formatida emas")


def _xlsx(content: bytes) -> list[Table]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [Table(ws.title, ws.iter_rows(values_only=True)) for ws in wb.worksheets]


def _xls(content: bytes) -> list[Table]:
    import xlrd
    book = xlrd.open_workbook(file_contents=content)
    return [Table(sh.name, (tuple(sh.row_values(i)) for i in range(sh.nrows))) for sh in book.sheets()]


def _decode(content: bytes) -> str:
    # Excel'dan saqlangan CSV ko'pincha Windows-1251 (kirill) bo'ladi
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise invalid("bad_file", "CSV kodirovkasi aniqlanmadi: UTF-8 yoki Windows-1251 da saqlang")


def _csv(content: bytes) -> list[Table]:
    text = _decode(content)
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [Table("CSV", (tuple(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)))]


def _docx(content: bytes) -> list[Table]:
    from docx import Document
    doc = Document(io.BytesIO(content))

    def rows() -> Iterator[tuple]:
        # Hujjatdagi barcha jadvallar bitta ro'yxat: jadval sahifadan sahifaga bo'linib ketgan bo'lishi mumkin
        for table in doc.tables:
            for row in table.rows:
                yield tuple(cell.text for cell in row.cells)
    return [Table("Word jadvali", list(rows()))]


def _pdf(content: bytes) -> list[Table]:
    import pdfplumber
    rows: list[tuple] = []
    has_text = False
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            if (page.extract_text() or "").strip():
                has_text = True
            for table in page.extract_tables():
                rows += [tuple(c or "" for c in r) for r in table]
    if not has_text:
        raise invalid("pdf_scanned", "Bu PDF skanerlangan (ichida matn yo'q). Matnni o'qish (OCR) keyingi bosqichda qo'shiladi — "
                                     "hozircha Excel yoki Word faylini yuklang")
    return [Table("PDF jadvali", rows)]


READERS = {"xlsx": _xlsx, "xls": _xls, "csv": _csv, "docx": _docx, "pdf": _pdf}


def read_tables(file_name: str, content: bytes) -> tuple[str, list[Table]]:
    kind = file_kind(file_name)
    try:
        return kind, READERS[kind](content)
    except Exception as e:  # noqa: BLE001 — buzilgan fayl har xil xato beradi
        if getattr(e, "code", None):  # bizning AppError (pdf_scanned, bad_file)
            raise
        raise _bad(kind) from e
