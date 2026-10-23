"""Call-center Excel faylidan korxonalarni import qilish.

1. Sarlavha qatori bo'yicha varaq topiladi ("Ташкилот номи" + "Туман"); birinchi mos varaq olinadi —
   joriy faylda bu "Рес бўйича" (qolgan varaqlar uning tumanlar bo'yicha nusxasi).
2. Viloyat/tuman nomlari normallashtiriladi va bazadagi tumanga bog'lanadi; topilmasa — qator xato ro'yxatiga.
3. Qatorlar "tuman + normallashtirilgan nom" bo'yicha bitta korxonaga birlashtiriladi; barcha telefonlar saqlanadi.
4. Bazada shu kalit bilan korxona bo'lsa — yangi aloqalar qo'shiladi, bo'sh maydonlar to'ldiriladi (qayta import xavfsiz).
Buyurtmalar import qilinmaydi (qaror: faqat korxonalar).
"""
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.errors import invalid
from app.models import Customer, CustomerContact, District, ImportBatch, Region, User
from app.models.enums import LocationStatus
from app.services.normalize import format_phone, is_city, key, phone_digits, place_key

MAX_ROWS = 50_000
MAX_ERRORS_IN_REPORT = 300

# Sarlavha → maydon (normallashtirilgan sarlavhada shu bo'lak bo'lsa)
COLUMNS = {
    "region": ("hudud nomi", "viloyat"),
    "district": ("tuman shahar nomi", "tuman"),
    "name": ("tashkilot nomi",),
    "activity": ("faoliyat turi",),
    "address": ("manzili", "manzil"),
    "phone": ("telefon",),
    "contact": ("buyurtma bergan xodim",),
    "info": ("buyurtma haqida",),
    "stir": ("inn", "stir"),
}
SEPTIC = re.compile(r"assenizator|asenizator|assinizator|asinizator")


@dataclass
class RowError:
    row: int
    message: str


@dataclass
class Group:
    district_id: int
    name_key: str
    names: Counter = field(default_factory=Counter)
    activities: Counter = field(default_factory=Counter)
    addresses: Counter = field(default_factory=Counter)
    contacts: dict[str, str | None] = field(default_factory=dict)  # telefon raqamlari (9 xona) → ism
    stir: str | None = None
    rows: int = 0


@dataclass
class Plan:
    sheet: str
    rows_total: int = 0
    rows_ok: int = 0
    septic_rows: int = 0
    errors: list[RowError] = field(default_factory=list)
    groups: dict[tuple[int, str], Group] = field(default_factory=dict)


def _header_map(row: tuple) -> dict[str, int] | None:
    keys = [key(str(c)) if c is not None else "" for c in row]
    found: dict[str, int] = {}
    for col, needles in COLUMNS.items():
        for i, k in enumerate(keys):
            if i not in found.values() and any(n in k for n in needles):
                found[col] = i
                break
    return found if {"name", "district", "region"} <= found.keys() else None


def _cell(row: tuple, idx: int | None) -> str:
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return re.sub(r"\s+", " ", str(row[idx])).strip()


async def _district_index(db: AsyncSession) -> dict[tuple[str, str, bool], int]:
    rows = (await db.execute(select(District, Region).join(Region, District.region_id == Region.id).where(District.is_active))).all()
    return {(place_key(r.name), place_key(d.name), is_city(d.name)): d.id for d, r in rows}


async def build_plan(db: AsyncSession, content: bytes) -> Plan:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — har qanday buzilgan fayl
        raise invalid("bad_excel", "Faylni o'qib bo'lmadi: .xlsx formatidagi Excel fayl yuklang")
    districts = await _district_index(db)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header, head_row = None, 0
        for i, row in enumerate(rows, start=1):
            header = _header_map(row)
            if header or i >= 10:
                head_row = i
                break
        if not header:
            continue
        plan = Plan(sheet=ws.title)
        for n, row in enumerate(rows, start=head_row + 1):
            name = _cell(row, header.get("name"))
            if not name:
                continue  # bo'sh qator
            plan.rows_total += 1
            if plan.rows_total > MAX_ROWS:
                raise invalid("too_many_rows", f"Faylda {MAX_ROWS} dan ortiq qator bor")
            region, district = _cell(row, header["region"]), _cell(row, header["district"])
            district_id = districts.get((place_key(region), place_key(district), is_city(district)))
            if not district_id:
                plan.errors.append(RowError(n, f"Tuman topilmadi: «{region} — {district}»"))
                continue
            plan.rows_ok += 1
            if SEPTIC.search(key(_cell(row, header.get("info")))):
                plan.septic_rows += 1
            k = key(name)
            g = plan.groups.setdefault((district_id, k), Group(district_id=district_id, name_key=k))
            g.rows += 1
            g.names[name] += 1
            if a := _cell(row, header.get("activity")):
                g.activities[a] += 1
            if a := _cell(row, header.get("address")):
                g.addresses[a] += 1
            digits = phone_digits(_cell(row, header.get("phone")))
            if len(digits) == 9:
                g.contacts.setdefault(digits, _cell(row, header.get("contact")) or None)
            elif digits:
                plan.errors.append(RowError(n, f"Telefon raqami noto'g'ri, saqlanmadi: «{_cell(row, header.get('phone'))}»"))
            stir = re.sub(r"\D", "", _cell(row, header.get("stir")))
            if len(stir) == 9:
                g.stir = stir
        return plan
    raise invalid("no_sheet", "Faylda \"Ташкилот номи\", \"Ҳудуд номи\" va \"Туман\" ustunlari bor varaq topilmadi")


@dataclass
class ImportResult:
    sheet: str
    rows_total: int
    rows_ok: int
    rows_skipped: int
    septic_rows: int
    customers_new: int
    customers_existing: int
    contacts_added: int
    by_district: dict[str, int]
    errors: list[dict]
    batch_id: int | None = None


def _top(counter: Counter) -> str | None:
    return counter.most_common(1)[0][0] if counter else None


async def next_fixed_no_start(db: AsyncSession) -> int:
    """Oxirgi A-#### raqamidan keyingisi (qo'lda va importda bir xil ketma-ketlik)."""
    nums = [int(m.group(1)) for (f,) in (await db.execute(select(Customer.fixed_no))).all()
            if (m := re.fullmatch(r"A-(\d+)", f or ""))]
    return max(nums, default=100) + 1


async def run(db: AsyncSession, user: User, file_name: str, content: bytes, dry_run: bool) -> ImportResult:
    plan = await build_plan(db, content)
    existing = {(c.district_id, c.name_key): c for c in (await db.execute(
        select(Customer).options(selectinload(Customer.contacts)))).scalars()}
    stirs = {s for (s,) in (await db.execute(select(Customer.stir).where(Customer.stir.is_not(None)))).all()}
    district_names = {d.id: d.name for d in (await db.execute(select(District))).scalars()}
    next_no = await next_fixed_no_start(db)
    batch = None
    if not dry_run:
        batch = ImportBatch(file_name=file_name[:255], uploaded_by_id=user.id)
        db.add(batch)
        await db.flush()

    new = merged = contacts_added = 0
    by_district: Counter = Counter()
    for (district_id, name_key), g in plan.groups.items():
        by_district[district_names[district_id]] += 1
        stir = g.stir if g.stir and g.stir not in stirs else None
        cust = existing.get((district_id, name_key))
        if cust is None:
            new += 1
            if stir:
                stirs.add(stir)
            if dry_run:
                contacts_added += len(g.contacts)
                continue
            first_phone = next(iter(g.contacts), None)
            cust = Customer(fixed_no=f"A-{next_no:04d}", name=_top(g.names), name_key=name_key,
                            activity_type=_top(g.activities), address=_top(g.addresses), district_id=district_id, stir=stir,
                            phone=format_phone(first_phone), location_status=LocationStatus.none, source="excel",
                            import_batch_id=batch.id)
            next_no += 1
            cust.contacts = [CustomerContact(name=nm, phone=format_phone(ph)) for ph, nm in g.contacts.items()]
            db.add(cust)
            contacts_added += len(g.contacts)
        else:
            merged += 1
            known = {phone_digits(c.phone) for c in cust.contacts} | {phone_digits(cust.phone)}
            fresh = {ph: nm for ph, nm in g.contacts.items() if ph not in known}
            contacts_added += len(fresh)
            if dry_run:
                continue
            for ph, nm in fresh.items():
                cust.contacts.append(CustomerContact(name=nm, phone=format_phone(ph)))
            cust.activity_type = cust.activity_type or _top(g.activities)
            cust.address = cust.address or _top(g.addresses)
            cust.phone = cust.phone or format_phone(next(iter(g.contacts), None))
            if not cust.stir and stir:
                cust.stir = stir
                stirs.add(stir)

    errors = [{"row": e.row, "message": e.message} for e in plan.errors[:MAX_ERRORS_IN_REPORT]]
    result = ImportResult(sheet=plan.sheet, rows_total=plan.rows_total, rows_ok=plan.rows_ok,
                          rows_skipped=plan.rows_total - plan.rows_ok, septic_rows=plan.septic_rows, customers_new=new,
                          customers_existing=merged, contacts_added=contacts_added, by_district=dict(by_district), errors=errors)
    if dry_run:
        await db.rollback()
        return result
    batch.rows_total, batch.customers_created, batch.customers_merged, batch.rows_skipped = (
        plan.rows_total, new, merged, result.rows_skipped)
    batch.report = json.dumps({"sheet": plan.sheet, "septic_rows": plan.septic_rows, "contacts_added": contacts_added,
                               "by_district": result.by_district, "errors": errors}, ensure_ascii=False)
    await db.commit()
    result.batch_id = batch.id
    return result
