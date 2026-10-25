"""Statistikani Excel (.xlsx) va PDF ko'rinishida yuklab olish — tanlangan davr (va tanlangan ko'rsatkich) bo'yicha.

Fayl tarkibi: umumiy ko'rsatkichlar, tumanlar kesimi (barcha ko'rsatkichlar), tanlangan ko'rsatkichning dinamikasi,
tarkibi va bajarilish vaqti taqsimoti. PDF kirill/lotin uchun DejaVu shriftidan foydalanadi (app/assets/fonts).
"""
import io
from datetime import timedelta
from pathlib import Path

from app.schemas import StatsOut

PERIOD_LABEL = {"today": "Bugun", "month": "Shu oy", "last10": "Oxirgi 10 kun", "last5": "Oxirgi 5 kun"}
METRIC_LABEL = {"received": "Tushgan arizalar", "repeat": "Qayta arizalar", "pending": "Kutilmoqda (muddat ichida)",
                "overdue": "Muddatdan oshgan (72 soat)", "done": "Bajarilgan", "not_done": "Bajarilmagan"}
COMPOSITION_LABEL = {"on_time": "Muddatida", "late": "Kechikib", "suspect": "Shubhali", "unassigned": "Biriktirilmagan",
                     "in_progress": "Jarayonda", "failed": "Bajarib bo'lmadi", "done": "Bajarilgan", "pending": "Kutilmoqda",
                     "overdue": "Muddati o'tgan"}
COLUMNS = ["received", "repeat", "pending", "overdue", "done", "not_done"]
FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"


def _since_local(s: StatsOut) -> str:
    return (s.since + timedelta(hours=5)).strftime("%d.%m.%Y %H:%M")


def _district_rows(s: StatsOut) -> list[list]:
    rows = []
    for d in s.districts:
        rows.append([d.name, *[getattr(d, c) for c in COLUMNS], d.done_share, d.not_done_share])
    t = s.total
    rows.append(["Jami", *[getattr(t, c) for c in COLUMNS], t.done_share, t.not_done_share])
    return rows


HEAD = ["Tuman", *[METRIC_LABEL[c] for c in COLUMNS], "Bajarilgan, %", "Bajarilmagan, %"]


def to_xlsx(s: StatsOut, d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8F1EC")
    wb = Workbook()

    def sheet(ws, title_rows: list[list], header: list[str], rows: list[list], widths: list[int]):
        for r in title_rows:
            ws.append(r)
        ws.append([])
        ws.append(header)
        hr = ws.max_row
        for c in ws[hr]:
            c.font, c.fill = bold, fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    period = [["Davr", PERIOD_LABEL[s.period]], ["Boshlanishi", _since_local(s)]]
    ws = wb.active
    ws.title = "Tumanlar"
    sheet(ws, [["Arizalar statistikasi — tumanlar kesimida"], *period], HEAD, _district_rows(s), [22] + [14] * 8)
    ws["A1"].font = Font(bold=True, size=13)
    m = METRIC_LABEL[d["metric"]]
    sheet(wb.create_sheet("Dinamika"), [[f"{m} — dinamika"], *period], ["Soat" if d["granularity"] == "hour" else "Sana", m],
          [[p.get("date", p["label"]), p["value"]] for p in d["series"]], [14, 18])
    sheet(wb.create_sheet("Tarkib"), [[f"{m} — tarkibi"], *period], ["Toifa", "Soni"],
          [[COMPOSITION_LABEL[c["key"]], c["value"]] for c in d["composition"]], [22, 12])
    sheet(wb.create_sheet("Vaqt taqsimoti"), [[f"{m} — ochilgandan bajarilguncha (yoki hozirgacha) o'tgan vaqt"], *period],
          ["Soat", "Soni"], [[h["bucket"], h["value"]] for h in d["histogram"]], [14, 12])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(s: StatsOut, d: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
    h1 = ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="DejaVu-Bold", fontSize=11, leading=14, spaceBefore=10, spaceAfter=4)
    p = ParagraphStyle("p", fontName="DejaVu", fontSize=9, leading=12, textColor=colors.HexColor("#52514e"))
    cell = ParagraphStyle("c", fontName="DejaVu", fontSize=7.5, leading=9)
    headc = ParagraphStyle("hc", fontName="DejaVu-Bold", fontSize=7.5, leading=9)

    def table(header: list[str], rows: list[list], widths: list[float], total_row: bool = False) -> Table:
        data = [[Paragraph(str(x), headc) for x in header]] + [[Paragraph(str(x), cell) for x in r] for r in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c9cfc9")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8f1ec")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 1), (-1, -1), "RIGHT")]
        if total_row:
            style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f3f5f2")))
        t.setStyle(TableStyle(style))
        return t

    m = METRIC_LABEL[d["metric"]]
    story = [Paragraph("Arizalar statistikasi", h1),
             Paragraph(f"Davr: {PERIOD_LABEL[s.period]} ({_since_local(s)} dan beri). Tanlangan ko'rsatkich: {m}.", p),
             Paragraph("Tumanlar kesimida", h2),
             table(HEAD, _district_rows(s), [40 * mm] + [26 * mm] * 8, total_row=True),
             Paragraph(f"{m} — tarkibi", h2),
             table(["Toifa", "Soni"], [[COMPOSITION_LABEL[c["key"]], c["value"]] for c in d["composition"]], [60 * mm, 25 * mm]),
             Paragraph(f"{m} — ochilgandan bajarilguncha (yoki hozirgacha) o'tgan vaqt", h2),
             table(["Soat", "Soni"], [[h["bucket"], h["value"]] for h in d["histogram"]], [60 * mm, 25 * mm]),
             Paragraph(f"{m} — dinamika", h2),
             table(["Soat" if d["granularity"] == "hour" else "Sana", "Soni"],
                   [[x.get("date", x["label"]), x["value"]] for x in d["series"]], [60 * mm, 25 * mm]),
             Spacer(1, 6)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                      title="Arizalar statistikasi").build(story)
    return buf.getvalue()
