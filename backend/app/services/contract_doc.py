"""Shartnoma hujjati: shablon bloklaridan PDF (chop etish) va DOCX (Didox / hujjat.uz'ga yuklash) — mazmun bir xil."""
import io
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.order_form_pdf import _fonts


def _html(t: str) -> str:
    return escape(t).replace("\n", "<br/>")


def render_pdf(blocks: list, title: str, paper: bool, draft_mark: str | None = None) -> bytes:
    _fonts()
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVu", 8)
        canvas.drawCentredString(A4[0] / 2, 9 * mm, str(doc.page))
        if draft_mark:
            canvas.setFillColor(colors.Color(0.8, 0, 0, alpha=0.25))
            canvas.setFont("DejaVu-Bold", 9)
            canvas.drawString(15 * mm, A4[1] - 9 * mm, draft_mark)
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=16 * mm, title=title)
    p = ParagraphStyle("p", fontName="DejaVu", fontSize=9.5, leading=12.6, alignment=TA_JUSTIFY, spaceAfter=3)
    small = ParagraphStyle("s", parent=p, fontSize=7.2, leading=8.8, alignment=TA_CENTER, spaceAfter=0)
    h = ParagraphStyle("h", parent=p, fontName="DejaVu-Bold", alignment=TA_CENTER, spaceBefore=7, spaceAfter=4)
    t = ParagraphStyle("t", parent=h, fontSize=13, leading=16, spaceBefore=0)
    right = ParagraphStyle("r", parent=p, alignment=TA_RIGHT)
    story = []
    for b in blocks:
        kind = b[0]
        if kind == "title":
            story.append(Paragraph(_html(b[1]), t))
        elif kind == "pair":
            story += [Spacer(1, 2 * mm), Table([[Paragraph(_html(b[1]), p), Paragraph(_html(b[2]), right)]], colWidths=[88 * mm, 89 * mm],
                                              style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)])), Spacer(1, 3 * mm)]
        elif kind == "h":
            story.append(Paragraph(_html(b[1]), h))
        elif kind == "p":
            story.append(Paragraph(_html(b[1]), p))
        elif kind == "table":
            spec = b[1]
            w = spec["widths"]
            total = 177 * mm
            cols = [total * x / sum(w) for x in w]
            rows = [[Paragraph(_html(x), small) for x in spec["head"]]] + [[Paragraph(_html(str(x)), small) for x in r] for r in spec["rows"]]
            rows.append([Paragraph(f"<b>{_html(str(x))}</b>", small) for x in spec["foot"]])
            tbl = Table(rows, colWidths=cols, repeatRows=1)
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                     ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
            story += [Spacer(1, 2 * mm), tbl, Spacer(1, 2 * mm)]
        elif kind == "req":
            spec = b[1]
            def col(title, items):
                return Paragraph(f"<b>{_html(title)}</b><br/><br/>" + "<br/>".join(f"<b>{_html(k)}:</b> {_html(v or '')}" for k, v in items), p)
            story.append(KeepTogether([Spacer(1, 2 * mm), Table([[col(spec["left_title"], spec["left"]), col(spec["right_title"], spec["right"])]],
                                       colWidths=[88 * mm, 89 * mm], style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))]))
        elif kind == "sign":
            if paper:
                story += [Spacer(1, 10 * mm), Table([["Imzo: ____________________", "Imzo: ____________________"], ["M.O‘.", "M.O‘."]],
                                                    colWidths=[88 * mm, 89 * mm],
                                                    style=TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                                                                      ("TOPPADDING", (0, 1), (-1, 1), 6)]))]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _cell_text(cell, text: str, bold=False, size=7, center=True):
    cell.text = ""
    par = cell.paragraphs[0]
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(text)
    run.bold, run.font.size = bold, Pt(size)


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def render_docx(blocks: list, title: str, paper: bool) -> bytes:
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.5), Cm(1.5), Cm(1.6)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    st.paragraph_format.space_after = Pt(3)
    d.core_properties.title = title
    for b in blocks:
        kind = b[0]
        if kind == "title":
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(b[1])
            r.bold, r.font.size = True, Pt(14)
        elif kind == "pair":
            tbl = d.add_table(rows=1, cols=2)
            _no_borders(tbl)
            tbl.cell(0, 0).text = b[1]
            tbl.cell(0, 1).text = b[2]
            for par in tbl.cell(0, 1).paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif kind == "h":
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(8)
            par.add_run(b[1]).bold = True
        elif kind == "p":
            for i, line in enumerate(b[1].split("\n")):
                par = d.add_paragraph(line)
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if i:
                    par.paragraph_format.space_before = Pt(0)
        elif kind == "table":
            spec = b[1]
            n = len(spec["head"])
            tbl = d.add_table(rows=2 + len(spec["rows"]), cols=n)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, x in enumerate(spec["head"]):
                _cell_text(tbl.cell(0, i), x, bold=True)
            for ri, r in enumerate(spec["rows"], start=1):
                for i, x in enumerate(r):
                    _cell_text(tbl.cell(ri, i), str(x))
            for i, x in enumerate(spec["foot"]):
                _cell_text(tbl.cell(len(spec["rows"]) + 1, i), str(x), bold=True)
        elif kind == "req":
            spec = b[1]
            tbl = d.add_table(rows=1, cols=2)
            _no_borders(tbl)
            for ci, (head, items) in enumerate(((spec["left_title"], spec["left"]), (spec["right_title"], spec["right"]))):
                cell = tbl.cell(0, ci)
                cell.text = ""
                cell.paragraphs[0].add_run(head).bold = True
                for k, v in items:
                    par = cell.add_paragraph()
                    par.add_run(f"{k}: ").bold = True
                    par.add_run(v or "")
        elif kind == "sign" and paper:
            d.add_paragraph()
            tbl = d.add_table(rows=2, cols=2)
            _no_borders(tbl)
            for ci in range(2):
                tbl.cell(0, ci).text = "Imzo: ____________________"
                tbl.cell(1, ci).text = "M.O‘."
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
